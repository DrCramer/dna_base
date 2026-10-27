from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.protocols import protocol_object_count


THIN = Side(style="thin", color="000000")


def _plate_sheet(workbook: Workbook, title: str, plate: dict[str, Any]) -> None:
    ws = workbook.create_sheet(title[:31])
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.append([title])
    ws.merge_cells("A1:M1")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append([""] + list(range(1, 13)))
    by_well = {item["well"]: item for item in plate["wells"]}
    for row_name in "ABCDEFGH":
        ws.append([row_name] + [by_well[f"{row_name}{column}"].get("display_name") or "" for column in range(1, 13)])
    for row in ws.iter_rows(min_row=2, max_row=10, min_col=1, max_col=13):
        for cell in row:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column in range(2, 14):
        ws.column_dimensions[get_column_letter(column)].width = 14
    for row in range(3, 11):
        ws.row_dimensions[row].height = 36
    ws.print_area = "A1:M10"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


def build_protocol_workbook(snapshot: dict[str, Any]) -> bytes:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Единая_плашка_A4"
    overview.sheet_view.showGridLines = False
    overview.page_setup.orientation = "landscape"
    overview.page_setup.paperSize = overview.PAPERSIZE_A4
    protocol = snapshot["protocol"]
    overview.append(["ЕДИНЫЙ ПРОТОКОЛ ПЛАШКИ: ВЫДЕЛЕНИЕ / RT / PCR / ФОРЕЗ"])
    overview.merge_cells("A1:H1")
    overview["A1"].font = Font(bold=True, size=15)
    overview["A1"].alignment = Alignment(horizontal="center")
    overview.append(["Дата", protocol["protocol_date"], "№", protocol["protocol_no"], "Название", protocol["name"], "Объектов", protocol_object_count(snapshot["objects"])])
    labels = {"dna_extraction": "Выделение", "realtime": "Real Time", "pcr": "PCR", "electrophoresis": "Форез"}
    for stage in snapshot["stages"]:
        names = ", ".join(item["display_name"] for item in stage.get("performers", []))
        overview.append([labels[stage["stage_type"]], stage.get("work_date"), stage.get("kit_name"), names, stage.get("sequencer_name"), stage.get("comment")])
    for row in overview.iter_rows(min_row=2, max_row=6, min_col=1, max_col=8):
        for cell in row:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    target_row = overview.max_row + 1
    overview.append(["Конечная концентрация ДНК (цель), нг/мкл", *[None] * 6, (snapshot.get("dilution_settings") or {}).get("target_concentration", 0.1)])
    overview.merge_cells(start_row=target_row, start_column=1, end_row=target_row, end_column=7)
    for cell in overview[target_row]:
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for key, title in (("source", "Исходная плашка"), ("pcr", "PCR")):
        for plate in snapshot["layouts"].get(key, {}).get("plates", []):
            _plate_sheet(workbook, f"{title} {plate['plate_index']}", plate)

    calculations = workbook.create_sheet("Расчёты")
    calculations.append(["Этап", "Плашка", "Компонент", "На реакцию", "Всего"])
    for block in snapshot["calculations"].get("pcr", []):
        for component in block["components"]:
            calculations.append(["PCR", block["plate_index"], component["label"], component["per_reaction"], None if component.get("key") == "dna" else component["total"]])
    for component in snapshot["calculations"].get("electrophoresis", {}).get("components", []):
        calculations.append(["Форез", "", component["label"], component["per_reaction"], None if component.get("key") == "pcr_products" else component["total"]])
    for cell in calculations[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    dilution = workbook.create_sheet("Dilution")
    optimization = snapshot.get("dilution_optimization") or {}
    dilution_headers = ["Плашка", "Лунка", "Объект", "Исх. конц.", "Кон. конц.", "Фактор", "Фактор I", "ДНК I", "H₂O I", "Фактор II", "ДНК II", "H₂O II"]
    show_targets = bool((snapshot.get("dilution_settings") or {}).get("show_all_targets") or (snapshot.get("dilution_settings") or {}).get("target_overrides"))
    rt_by_object = {obj.get("id"): obj.get("rt") or {} for obj in snapshot["objects"]}
    if show_targets:
        dilution_headers += ["Мишень", "Long", "Small", "Y", "Источник RT"]
    if optimization.get("enabled"):
        dilution_headers += ["Фактическая конц.", "Отклонение, %", "Базовая ДНК I", "Базовая H₂O I", "Базовая ДНК II", "Базовая H₂O II", "Базовая факт. конц.", "Статус округления", "Причина / примечание", "Остаток I разведения"]
    dilution_headers += ["Примечание"]
    dilution.append(dilution_headers)
    for item in snapshot.get("dilutions", []):
        steps = item.get("steps", [])
        first = steps[0] if steps else {}
        second = steps[1] if len(steps) > 1 else {}
        final_concentration = item.get("target_concentration")
        note = None
        if item.get("available") is False:
            final_concentration = None
        elif item.get("available") and not steps:
            final_concentration = item.get("source_concentration")
            note = "Без разведения"
            if final_concentration is not None and final_concentration < item["target_concentration"]:
                note += f": исходная концентрация ниже цели {item['target_concentration']:g} нг/мкл"
        values = [item.get("plate_index"), item.get("well"), item.get("display_name"), item.get("source_concentration"), final_concentration, item.get("total_factor"), first.get("factor"), first.get("dna_volume"), first.get("water_volume"), second.get("factor"), second.get("dna_volume"), second.get("water_volume")]
        if show_targets:
            targets = item.get("rt_targets") or {}
            rt = rt_by_object.get(item.get("object_id"), {})
            values += [item.get("active_target", "small"), targets.get("long"), targets.get("small"), targets.get("y"), rt.get("run_id") or rt.get("event_id")]
        if optimization.get("enabled"):
            metadata = item.get("optimization") or {}
            baseline = item.get("baseline_steps") or []
            base_one = baseline[0] if baseline else {}
            base_two = baseline[1] if len(baseline) > 1 else {}
            status = {"optimized": "Оптимизировано", "unchanged": "Исходный расчёт", "blocked": "Требует проверки"}.get(metadata.get("status"))
            values += [metadata.get("actual_final_concentration"), metadata.get("relative_error_percent"), base_one.get("dna_volume"), base_one.get("water_volume"), base_two.get("dna_volume"), base_two.get("water_volume"), metadata.get("baseline_actual_final_concentration"), status, metadata.get("reason"), metadata.get("first_dilution_remaining_volume")]
        dilution.append([*values, note])
    for cell in dilution[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    if optimization.get("enabled"):
        settings = snapshot.get("dilution_settings") or {}
        summary = workbook.create_sheet("Округление")
        summary.append(["Допуск и параметры дозирования требуют утверждения лабораторией"])
        summary.append(["Допуск, %", settings.get("allowed_concentration_error_percent")])
        summary.append(["Максимум ниже цели, %", settings.get("max_under_target_pct") if settings.get("max_under_target_pct") is not None else settings.get("allowed_concentration_error_percent")])
        summary.append(["Максимум выше цели, %", settings.get("max_over_target_pct") if settings.get("max_over_target_pct") is not None else settings.get("allowed_concentration_error_percent")])
        summary.append(["Предпочтительный шаг воды, мкл", settings.get("preferred_water_step_ul")])
        summary.append(["Приоритет", settings.get("optimization_priority")])
        summary.append(["Минимальный объём включён", settings.get("minimum_final_volume_enabled")])
        summary.append(["Минимальная смесь, мкл", settings.get("minimum_final_volume")])
        summary.append(["Доступная исходная ДНК, мкл", settings.get("source_available_volume")])
        for key, value in (settings.get("pipetting") or {}).items():
            summary.append([key, value])
        for key, label in (("unique_water_volumes_before", "Вариантов H₂O до"), ("unique_water_volumes_after", "Вариантов H₂O после"), ("optimized_object_count", "Оптимизировано объектов"), ("unchanged_object_count", "Сохранён базовый расчёт"), ("blocked_object_count", "Требуют проверки"), ("excluded_object_count", "Без расчёта / разведение не нужно"), ("max_absolute_error_percent", "Максимальное отклонение, %")):
            summary.append([label, optimization.get(key)])
        groups = workbook.create_sheet("Группы H₂O")
        groups.append(["Этап", "H₂O, мкл", "Плашка", "Лунка", "Объект", "Требует проверки"])
        for key, stage in (("batch_groups_first", "I"), ("batch_groups_second", "II")):
            for batch in optimization.get(key) or []:
                for item in batch["objects"]:
                    groups.append([stage, batch["water_volume"], item.get("plate_index"), item.get("well"), item.get("display_name"), bool(item.get("needs_review"))])
        for cell in groups[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for ws in workbook.worksheets:
        ws.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
