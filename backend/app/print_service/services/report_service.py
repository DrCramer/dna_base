from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


CSV_COLUMNS = [
    "Группа",
    "Столбец документов",
    "Столбец меток",
    "Порядок",
    "Номер документа",
    "Исходный файл",
    "Наносимая метка",
    "Метка нанесена",
    "Страница документа",
    "Страница итогового PDF",
    "Итоговый PDF",
    "Статус сопоставления",
    "Статус конвертации",
    "Количество страниц",
    "Ширина страницы, мм",
    "Высота страницы, мм",
    "Предупреждения",
    "Ошибка",
]


def write_csv_report(entries: list[dict[str, Any]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_COLUMNS, delimiter=";")
        writer.writeheader()
        for entry in entries:
            width = ""
            height = ""
            analysis = entry.get("pdf_analysis")
            if analysis and analysis.get("pages"):
                width = analysis["pages"][0]["width_mm"]
                height = analysis["pages"][0]["height_mm"]
            writer.writerow(
                {
                    "Порядок": entry.get("order", ""),
                    "Группа": entry.get("group", ""),
                    "Столбец документов": entry.get("group_column", ""),
                    "Столбец меток": entry.get("stamp_column", ""),
                    "Номер документа": entry.get("number", ""),
                    "Исходный файл": entry.get("matched_file") or "",
                    "Наносимая метка": entry.get("stamp_label") or "",
                    "Метка нанесена": "да" if entry.get("stamp_applied") else "нет",
                    "Страница документа": entry.get("pages") or "",
                    "Страница итогового PDF": entry.get("final_page") or entry.get("order", ""),
                    "Итоговый PDF": entry.get("result_pdf_name") or "",
                    "Статус сопоставления": entry.get("status", ""),
                    "Статус конвертации": entry.get("conversion_status", ""),
                    "Количество страниц": entry.get("pages") or "",
                    "Ширина страницы, мм": width,
                    "Высота страницы, мм": height,
                    "Предупреждения": " | ".join(entry.get("warnings") or []),
                    "Ошибка": entry.get("error") or "",
                }
            )


def write_number_mapping_xlsx(entries: list[dict[str, Any]], output_path: Path) -> int:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сопоставление"
    sheet.freeze_panes = "A2"
    sheet.append(["Исходный номер", "Присвоенный номер"])

    header_fill = PatternFill(fill_type="solid", fgColor="DCEAEC")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="203238")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for entry in entries:
        source_number = entry.get("external_military_no") or entry.get("number") or ""
        assigned_number = (
            entry.get("stamp_label")
            or entry.get("rcsme_reg_no")
            or entry.get("decree_no")
            or ""
        )
        sheet.append([str(source_number), str(assigned_number)])

    for row in sheet.iter_rows(min_row=2, min_col=1, max_col=2):
        for cell in row:
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top")
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 28
    sheet.auto_filter.ref = f"A1:B{max(1, sheet.max_row)}"
    workbook.save(output_path)
    workbook.close()
    return len(entries)
