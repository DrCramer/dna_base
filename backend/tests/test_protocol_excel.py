from io import BytesIO
import copy

from openpyxl import load_workbook

from app.services.protocol_excel import build_protocol_workbook
from app.services.protocol_plate import build_protocol_layouts
from app.services.protocol_calculations import calculate_dilution


def test_protocol_excel_keeps_snapshot_order_and_multiple_pcr_plates():
    objects = [{"id": index, "rcsme_reg_no": f"ии{index}", "order_index": index - 1} for index in range(1, 97)]
    layouts = build_protocol_layouts(objects, {"nc_enabled": False})
    snapshot = {
        "protocol": {"protocol_date": "2026-09-21", "protocol_no": 1, "name": "21-09-26_1"},
        "objects": objects,
        "stages": [
            {"stage_type": stage, "work_date": "2026-09-21", "kit_name": None, "performers": [], "sequencer_name": None, "comment": None}
            for stage in ("dna_extraction", "realtime", "pcr", "electrophoresis")
        ],
        "layouts": layouts,
        "calculations": {"pcr": [], "electrophoresis": {"components": []}},
        "dilutions": [],
    }
    workbook = load_workbook(BytesIO(build_protocol_workbook(snapshot)), data_only=True)
    assert "Исходная плашка 1" in workbook.sheetnames
    assert "PCR 1" in workbook.sheetnames
    assert "PCR 2" in workbook.sheetnames
    assert workbook["Исходная плашка 1"]["B3"].value == "ии1"
    assert workbook["PCR 1"]["C3"].value == "ии8"
    assert workbook["PCR 2"]["B4"].value == "ии89"


def test_excel_omits_legacy_sample_totals_without_changing_snapshot():
    objects = [{"id": 1, "rcsme_reg_no": "ии1"}]
    snapshot = {
        "protocol": {"protocol_date": "2026-10-14", "protocol_no": 1, "name": "test"},
        "objects": objects, "stages": [], "layouts": build_protocol_layouts(objects),
        "calculations": {
            "pcr": [{"plate_index": 1, "components": [
                {"key": "master_mix", "label": "Master mix", "per_reaction": 5, "total": 16.5},
                {"key": "dna", "label": "DNA", "per_reaction": 1, "total": 3.3},
            ]}],
            "electrophoresis": {"components": [
                {"key": "ils", "label": "ILS", "per_reaction": 0.5, "total": 4.4},
                {"key": "pcr_products", "label": "PCR-продукты", "per_reaction": 1, "total": 8.8},
            ]},
        },
        "dilutions": [],
    }
    original = copy.deepcopy(snapshot)
    workbook = load_workbook(BytesIO(build_protocol_workbook(snapshot)), data_only=True)
    rows = list(workbook["Расчёты"].values)[1:]
    assert rows == [
        ("PCR", 1, "Master mix", 5, 16.5), ("PCR", 1, "DNA", 1, None),
        ("Форез", None, "ILS", 0.5, 4.4), ("Форез", None, "PCR-продукты", 1, None),
    ]
    assert snapshot == original


def test_excel_does_not_claim_an_unreachable_target_or_missing_concentration():
    objects = [{"id": index, "rcsme_reg_no": str(index)} for index in range(1, 4)]
    snapshot = {"protocol": {"protocol_date": "2026-10-27", "protocol_no": 1, "name": "Target"}, "objects": objects, "stages": [], "layouts": build_protocol_layouts(objects), "calculations": {"pcr": [], "electrophoresis": {"components": []}}, "dilutions": [{"object_id": index, **calculate_dilution(source, target_concentration=1)} for index, source in enumerate([0.4, None, 10], 1)]}
    snapshot["dilution_settings"] = {"target_concentration": 1}
    original = copy.deepcopy(snapshot)
    workbook = load_workbook(BytesIO(build_protocol_workbook(snapshot)), data_only=True)
    sheet = workbook["Dilution"]
    assert workbook["Единая_плашка_A4"]["H7"].value == 1
    assert [sheet[f"E{row}"].value for row in (2, 3, 4)] == [0.4, None, 1]
    assert "ниже цели 1" in sheet["M2"].value
    assert sheet["H2"].value is None and sheet["I2"].value is None
    assert snapshot == original
