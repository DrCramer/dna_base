from pathlib import Path

from openpyxl import Workbook

from app.print_service.services import excel_service
from app.print_service.services.excel_service import (
    match_excel_files,
    match_excel_groups,
    parse_excel_sequences,
    rematch_excel_groups,
)


def docs(*names):
    return [
        {"id": f"doc_{index}", "original_name": name, "path": f"input/doc_{index}.docx"}
        for index, name in enumerate(names, start=1)
    ]


def write_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Первая пачка"
    sheet["A2"] = "ее5968"
    sheet["A3"] = "ее6032"
    sheet["B1"] = "Вторая пачка"
    sheet["B2"] = "нн2832"
    workbook.save(path)


def test_parse_excel_columns_as_separate_sequences(tmp_path):
    path = tmp_path / "order.xlsx"
    write_workbook(path)
    groups = parse_excel_sequences(
        path,
        docs("Документ_ее5968.docx", "Документ_ее6032.docx", "Документ_нн2832.docx"),
    )
    assert [group["title"] for group in groups] == ["Первая пачка", "Вторая пачка"]
    assert groups[0]["sequence"] == "ее5968\nее6032"
    assert groups[1]["sequence"] == "нн2832"


def test_parse_legacy_xls_uses_existing_excel_reader(tmp_path, monkeypatch):
    path = tmp_path / "legacy.xls"
    path.write_bytes(b"legacy")
    monkeypatch.setattr(
        excel_service,
        "read_workbook",
        lambda _path: {"Лист1": [["ии10"], ["ии2"]]},
    )

    groups = parse_excel_sequences(path, docs("Акт_ии10.docx", "Акт_ии2.docx"))

    assert groups[0]["sheet"] == "Лист1"
    assert groups[0]["sequence"] == "ии10\nии2"


def test_match_excel_groups_reports_global_unused_documents(tmp_path):
    path = tmp_path / "order.xlsx"
    write_workbook(path)
    result = match_excel_groups(
        path,
        docs(
            "Документ_ее5968.docx",
            "Документ_ее6032.docx",
            "Документ_нн2832.docx",
            "Лишний_ии0001.docx",
        ),
    )
    assert result["mode"] == "excel"
    assert result["can_build"] is True
    assert result["total_groups"] == 2
    assert len(result["unused_documents"]) == 1


def test_multiple_excel_files_keep_same_columns_as_distinct_groups(tmp_path):
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    for path, values in ((first, ["ии10", "ии2"]), (second, ["ии100", "ии1"])):
        workbook = Workbook()
        sheet = workbook.active
        for row, value in enumerate(values, start=1):
            sheet.cell(row=row, column=1, value=value)
        workbook.save(path)
        workbook.close()

    result = match_excel_files(
        [
            {"id": "excel_001", "name": "first.xlsx", "path": first},
            {"id": "excel_002", "name": "second.xlsx", "path": second},
        ],
        docs("Акт_ии1.docx", "Акт_ии2.docx", "Акт_ии10.docx", "Акт_ии100.docx"),
    )

    assert result["total_groups"] == 2
    assert [group["column"] for group in result["groups"]] == ["A", "A"]
    assert [group["excel_file"] for group in result["groups"]] == ["first.xlsx", "second.xlsx"]
    assert len({group["id"] for group in result["groups"]}) == 2
    assert result["groups"][0]["id"].startswith("excel_001:")
    assert result["groups"][1]["id"].startswith("excel_002:")


def test_excel_sorting_is_applied_independently_inside_each_group(tmp_path):
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    for path, values in ((first, ["ии10", "ии2", "ии1"]), (second, ["нн10", "нн2"])):
        workbook = Workbook()
        sheet = workbook.active
        for row, value in enumerate(values, start=1):
            sheet.cell(row=row, column=1, value=value)
        workbook.save(path)
        workbook.close()
    documents = docs(
        "Акт_ии1.docx",
        "Акт_ии2.docx",
        "Акт_ии10.docx",
        "Акт_нн2.docx",
        "Акт_нн10.docx",
    )
    validation = match_excel_files(
        [
            {"id": "excel_001", "name": "first.xlsx", "path": first},
            {"id": "excel_002", "name": "second.xlsx", "path": second},
        ],
        documents,
    )
    ordered = {
        validation["groups"][0]["id"]: ["ии1", "ии2", "ии10"],
        validation["groups"][1]["id"]: ["нн2", "нн10"],
    }

    result = rematch_excel_groups(validation, ordered, documents)

    assert [entry["number"] for entry in result["groups"][0]["validation"]["entries"]] == [
        "ии1",
        "ии2",
        "ии10",
    ]
    assert [entry["number"] for entry in result["groups"][1]["validation"]["entries"]] == [
        "нн2",
        "нн10",
    ]
