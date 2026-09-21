from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile
from io import BytesIO

import fitz
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pypdf import PdfWriter

from app.print_service import main
from app.print_service.services.job_store import get_job_dir, save_state
from app.print_service.services.report_service import write_number_mapping_xlsx


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(main.settings, "data_dir", tmp_path)
    app = FastAPI()
    app.include_router(main.router)
    app.include_router(main.api_router)
    app.dependency_overrides[main.current_user] = lambda: SimpleNamespace(id=1, username="viewer", role="viewer")
    with TestClient(app) as test_client:
        yield test_client


def test_print_api_requires_dna_session():
    app = FastAPI()
    app.include_router(main.api_router)

    with TestClient(app) as test_client:
        response = test_client.get("/api/print/jobs/recent")

    assert response.status_code == 401


def upload_docx(client, name="Акт_ее5968.docx"):
    return client.post(
        "/api/print/jobs",
        files=[("files", (name, b"fake-docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )


def test_health(client):
    response = client.get("/api/print/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_upload_without_docx_is_rejected(client):
    response = client.post("/api/print/jobs", files=[("files", ("list.txt", b"ee5968", "text/plain"))])
    assert response.status_code == 400


def test_job_upload_persists_external_military_excel(client):
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "ии8828"
    sheet["A2"] = "ии6305"
    buffer = BytesIO()
    workbook.save(buffer)

    response = client.post(
        "/api/print/jobs",
        files=[
            ("files", ("first.docx", b"docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("files", ("общий список.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
        ],
    )

    assert response.status_code == 200
    assert response.json()["registration_external_excel"]["all_labels"] == ["ии8828", "ии6305"]
    state = main.load_state(response.json()["id"])
    assert state["registration_external_numbers"] == ["ии8828", "ии6305"]


def test_validate_matching_error(client):
    upload = upload_docx(client, "Акт_ее5968.docx")
    job_id = upload.json()["id"]
    response = client.post(f"/api/print/jobs/{job_id}/validate", json={"sequence": "ее6032"})
    assert response.status_code == 200
    assert response.json()["can_build"] is False


def test_full_success_flow_with_mocked_conversion(client, monkeypatch):
    async def fake_convert_entries(entries, documents, job_dir, settings, progress_callback=None):
        pdf_paths = []
        for entry in entries:
            path = Path(job_dir) / "converted" / f"{entry['doc_id']}.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=595, height=842)
            with path.open("wb") as handle:
                writer.write(handle)
            entry["conversion_status"] = "Готов"
            entry["pages"] = 1
            entry["page_size"] = "210.0 × 297.0 мм · A4, книжная"
            entry["pdf_path"] = str(path.relative_to(job_dir))
            entry["pdf_analysis"] = {
                "page_count": 1,
                "pages": [{"width_mm": 210.0, "height_mm": 297.0, "rotate": 0}],
                "empty_pages": [],
            }
            pdf_paths.append(path)
        return entries, pdf_paths

    monkeypatch.setattr(main, "convert_entries", fake_convert_entries)
    upload = upload_docx(client)
    assert upload.status_code == 200
    job_id = upload.json()["id"]
    validation = client.post(f"/api/print/jobs/{job_id}/validate", json={"sequence": "ее5968"})
    assert validation.json()["can_build"] is True
    build = client.post(f"/api/print/jobs/{job_id}/build?wait=true")
    assert build.status_code == 200
    assert build.json()["merge"]["page_count"] == 1
    assert build.json()["result_pdfs"][0]["download_name"] == "1-(ее5968).pdf"
    state = main.load_state(job_id)
    assert state["result_pdf"] == "result/1-(ее5968).pdf"
    assert state["validation"]["entries"][0]["result_pdf_name"] == "1-(ее5968).pdf"
    download = client.get(f"/api/print/jobs/{job_id}/download/pdf")
    assert download.status_code == 200
    assert "%D0%B5%D0%B55968" in download.headers["content-disposition"]
    report = (get_job_dir(job_id) / state["report_csv"]).read_text(encoding="utf-8-sig")
    assert "1-(ее5968).pdf" in report
    assert client.get(f"/api/print/jobs/{job_id}/download/report.csv").status_code == 200
    mapping_download = client.get(f"/api/print/jobs/{job_id}/download/number-mapping.xlsx")
    assert mapping_download.status_code == 200
    assert "%D0%A1%D0%BE%D0%BF%D0%BE%D1%81%D1%82%D0%B0%D0%B2%D0%BB%D0%B5%D0%BD%D0%B8%D0%B5" in mapping_download.headers["content-disposition"]
    assert client.delete(f"/api/print/jobs/{job_id}").status_code == 200


def test_registration_build_creates_party_named_pdfs(client, monkeypatch):
    async def fake_convert_entries(entries, documents, job_dir, settings, progress_callback=None):
        pdf_paths = []
        for entry in entries:
            path = Path(job_dir) / "converted" / f"{entry['doc_id']}.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=595, height=842)
            with path.open("wb") as handle:
                writer.write(handle)
            entry["conversion_status"] = "Готов"
            entry["pages"] = 1
            entry["page_size"] = "210.0 × 297.0 мм · A4, книжная"
            entry["pdf_path"] = str(path.relative_to(job_dir))
            entry["pdf_analysis"] = {
                "page_count": 1,
                "pages": [{"width_mm": 210.0, "height_mm": 297.0, "rotate": 0}],
                "empty_pages": [],
            }
            pdf_paths.append(path)
        return entries, pdf_paths

    monkeypatch.setattr(main, "convert_entries", fake_convert_entries)
    upload = client.post(
        "/api/print/jobs",
        files=[
            ("files", ("first.docx", b"first", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("files", ("last.docx", b"last", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ],
    )
    job_id = upload.json()["id"]
    documents = upload.json()["documents"]
    state = main.load_state(job_id)
    state["status"] = "validated"
    state["validation"] = {
        "mode": "registration",
        "can_build": True,
        "entries": [
            {
                "order": 1,
                "doc_id": documents[0]["id"],
                "party_no": "194",
                "external_military_no": "ии8828",
                "rcsme_reg_no": "6406-1",
                "status": "Готов",
                "blocking": False,
                "warnings": [],
                "error": "",
            },
            {
                "order": 2,
                "doc_id": documents[1]["id"],
                "party_no": "194",
                "external_military_no": "ии6305",
                "rcsme_reg_no": "6505-1",
                "status": "Готов",
                "blocking": False,
                "warnings": [],
                "error": "",
            },
        ],
        "stamping": {"config": {"enabled": False}},
    }
    main.save_state(job_id, state)

    build = client.post(f"/api/print/jobs/{job_id}/build?wait=true")

    assert build.status_code == 200
    assert build.json()["mode"] == "registration"
    assert build.json()["result_pdfs"][0]["download_name"] == "1-(ии8828-ии6305).pdf"
    state = main.load_state(job_id)
    zip_path = main.get_job_dir(job_id) / state["result_zip"]
    with ZipFile(zip_path) as archive:
        assert "1-(ии8828-ии6305).pdf" in archive.namelist()
        assert "Сопоставление номеров.xlsx" in archive.namelist()
        assert "report.csv" not in archive.namelist()
    assert state["report_csv"] is None


def test_excel_build_uses_range_names_in_files_zip_and_report(client, monkeypatch):
    async def fake_convert_entries(entries, documents, job_dir, settings, progress_callback=None):
        pdf_paths = []
        for entry in entries:
            path = Path(job_dir) / "converted" / f"{entry['doc_id']}.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=595, height=842)
            with path.open("wb") as handle:
                writer.write(handle)
            entry["conversion_status"] = "Готов"
            entry["pages"] = 1
            entry["page_size"] = "210.0 × 297.0 мм · A4, книжная"
            entry["pdf_path"] = str(path.relative_to(job_dir))
            entry["pdf_analysis"] = {
                "page_count": 1,
                "pages": [{"width_mm": 210.0, "height_mm": 297.0, "rotate": 0}],
                "empty_pages": [],
            }
            pdf_paths.append(path)
        return entries, pdf_paths

    monkeypatch.setattr(main, "convert_entries", fake_convert_entries)
    upload = client.post(
        "/api/print/jobs",
        files=[
            ("files", ("first.docx", b"first", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("files", ("last.docx", b"last", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ],
    )
    job_id = upload.json()["id"]
    documents = upload.json()["documents"]
    state = main.load_state(job_id)
    state["status"] = "validated"
    state["validation"] = {
        "mode": "excel",
        "can_build": True,
        "groups": [
            {
                "title": "Столбец A",
                "column": "A",
                "validation": {
                    "total": 1,
                    "entries": [
                        {
                            "order": 1,
                            "number": "ее5968",
                            "doc_id": documents[0]["id"],
                            "status": "Готов",
                            "blocking": False,
                            "warnings": [],
                            "error": "",
                        }
                    ],
                },
            },
            {
                "title": "Столбец B",
                "column": "B",
                "validation": {
                    "total": 1,
                    "entries": [
                        {
                            "order": 1,
                            "number": "ее6032",
                            "doc_id": documents[1]["id"],
                            "status": "Готов",
                            "blocking": False,
                            "warnings": [],
                            "error": "",
                        }
                    ],
                },
            },
        ],
        "stamping": {"config": {"enabled": False}},
    }
    main.save_state(job_id, state)

    build = client.post(f"/api/print/jobs/{job_id}/build?wait=true")

    assert build.status_code == 200
    assert [item["download_name"] for item in build.json()["result_pdfs"]] == [
        "1-(ее5968).pdf",
        "2-(ее6032).pdf",
    ]
    state = main.load_state(job_id)
    with ZipFile(main.get_job_dir(job_id) / state["result_zip"]) as archive:
        assert set(archive.namelist()) == {
            "1-(ее5968).pdf",
            "2-(ее6032).pdf",
            "report.csv",
            "Сопоставление номеров.xlsx",
        }
    report = (main.get_job_dir(job_id) / state["report_csv"]).read_text(encoding="utf-8-sig")
    assert "1-(ее5968).pdf" in report
    assert "2-(ее6032).pdf" in report
    workbook = load_workbook(main.get_job_dir(job_id) / state["number_mapping_xlsx"], read_only=True, data_only=True)
    assert list(workbook["Сопоставление"].iter_rows(values_only=True))[1:] == [
        ("ее5968", None),
        ("ее6032", None),
    ]
    workbook.close()


@pytest.mark.parametrize(
    ("index", "entries", "expected"),
    [
        (1, [{"number": "ее5968"}], "1-(ее5968).pdf"),
        (1, [{"number": "ее5968"}, {"number": "ее6032"}], "1-(ее5968-ее6032).pdf"),
        (2, [{"number": "6606-1"}, {"number": "6610-1"}], "2-(6606-1-6610-1).pdf"),
        (3, [], "3-(без-диапазона).pdf"),
    ],
)
def test_result_pdf_name_uses_actual_entry_order(index, entries, expected):
    assert main._result_pdf_name(index, entries) == expected


def test_number_mapping_xlsx_preserves_final_pdf_order(tmp_path):
    path = tmp_path / "mapping.xlsx"
    entries = [
        {
            "number": "ии10/1",
            "source_number_original": "ии10/1",
            "assigned_number": "7701-2026",
            "matched_docx": "Акт_ии10_1.docx",
            "pdf_name": "1-(ии10_1-ии100).pdf",
            "pdf_position": 1,
        },
        {"number": "ии2", "stamp_label": "7700-2026"},
        {"number": "ии100", "stamp_label": "7702-2026"},
    ]

    assert write_number_mapping_xlsx(entries, path) == 3

    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    detail_rows = list(workbook["Детали"].iter_rows(values_only=True))
    workbook.close()
    assert rows == [
        ("Исходный номер", "Присвоенный номер"),
        ("ии10/1", "7701-2026"),
        ("ии2", "7700-2026"),
        ("ии100", "7702-2026"),
    ]
    assert detail_rows[0] == (
        "Исходный номер",
        "Присвоенный номер",
        "Файл DOCX",
        "PDF",
        "Позиция в PDF",
    )
    assert detail_rows[1] == (
        "ии10/1",
        "7701-2026",
        "Акт_ии10_1.docx",
        "1-(ии10_1-ии100).pdf",
        1,
    )


def test_upload_keeps_safe_relative_folder_path(client):
    response = client.post(
        "/api/print/jobs",
        files=[
            (
                "files",
                (
                    "Партия 200/Группа 1/Акт_ее5968.docx",
                    b"docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            )
        ],
    )

    assert response.status_code == 200
    document = response.json()["documents"][0]
    assert document["original_name"] == "Акт_ее5968.docx"
    assert document["source_relative_path"] == "Партия 200/Группа 1/Акт_ее5968.docx"


def test_full_success_flow_with_stamping(client, monkeypatch):
    async def fake_convert_entries(entries, documents, job_dir, settings, progress_callback=None):
        pdf_paths = []
        for entry in entries:
            path = Path(job_dir) / "converted" / f"{entry['doc_id']}.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=595, height=842)
            with path.open("wb") as handle:
                writer.write(handle)
            entry["conversion_status"] = "Готов"
            entry["pages"] = 1
            entry["page_size"] = "210.0 × 297.0 мм · A4, книжная"
            entry["pdf_path"] = str(path.relative_to(job_dir))
            entry["pdf_analysis"] = {
                "page_count": 1,
                "pages": [{"width_mm": 210.0, "height_mm": 297.0, "rotate": 0}],
                "empty_pages": [],
            }
            pdf_paths.append(path)
        return entries, pdf_paths

    monkeypatch.setattr(main, "convert_entries", fake_convert_entries)
    upload = upload_docx(client)
    job_id = upload.json()["id"]
    validation = client.post(
        f"/api/print/jobs/{job_id}/validate",
        json={
            "sequence": "ее5968",
            "stamping": {"enabled": True, "text": "6528-2026", "style": {}},
        },
    )
    assert validation.status_code == 200
    assert validation.json()["can_build"] is True
    build = client.post(f"/api/print/jobs/{job_id}/build?wait=true")
    assert build.status_code == 200
    assert build.json()["stamping"]["applied"] == 1

    state = main.load_state(job_id)
    result_pdf = get_job_dir(job_id) / state["result_pdf"]
    document = fitz.open(result_pdf)
    try:
        assert "6528-2026" in document[0].get_text("text")
    finally:
        document.close()
    report = (get_job_dir(job_id) / state["report_csv"]).read_text(encoding="utf-8-sig")
    assert "Наносимая метка" in report
    assert "6528-2026" in report
    workbook = load_workbook(get_job_dir(job_id) / state["number_mapping_xlsx"], read_only=True, data_only=True)
    assert list(workbook.active.iter_rows(values_only=True))[1] == ("ее5968", "6528-2026")
    workbook.close()


def test_stamping_count_mismatch_blocks_build(client):
    upload = upload_docx(client)
    job_id = upload.json()["id"]
    response = client.post(
        f"/api/print/jobs/{job_id}/validate",
        json={
            "sequence": "ее5968\nее5968",
            "stamping": {"enabled": True, "text": "6528-2026", "style": {}},
        },
    )
    assert response.status_code == 200
    assert response.json()["can_build"] is False
    assert "меток" in response.json()["blocking_errors"][-1]


def test_stamping_preview_before_validation_does_not_save_validation(client, monkeypatch):
    def fake_convert_docx_to_pdf(doc, job_dir, settings):
        path = Path(job_dir) / "converted" / f"{doc['id']}.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        with path.open("wb") as handle:
            writer.write(handle)
        return path

    monkeypatch.setattr(main, "convert_docx_to_pdf", fake_convert_docx_to_pdf)
    upload = upload_docx(client)
    job_id = upload.json()["id"]

    response = client.post(
        f"/api/print/jobs/{job_id}/stamping/preview",
        json={
            "sequence": "ее5968",
            "stamping": {"enabled": True, "text": "6528-2026", "style": {}},
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "image/png"
    state = main.load_state(job_id)
    assert state["validation"] is None
    assert state["base_validation"] is None


def test_excel_validation_ignores_enabled_stamping_without_group_labels(client, tmp_path):
    upload = upload_docx(client)
    job_id = upload.json()["id"]
    xlsx_path = tmp_path / "order.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "ее5968"
    workbook.save(xlsx_path)
    workbook.close()

    with xlsx_path.open("rb") as handle:
        response = client.post(
            f"/api/print/jobs/{job_id}/validate/excel",
            files=[
                (
                    "file",
                    ("order.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                )
            ],
            data={
                "stamping_json": (
                    '{"enabled": true, "groups": {}, "style": '
                    '{"corner": "top_left", "rotation": "left_90"}}'
                )
            },
        )

    assert response.status_code == 200
    data = response.json()
    assert data["can_build"] is True
    assert data["stamping"]["config"]["enabled"] is False
    assert data["stamping"]["summary"]["labels"] == 0


def test_excel_validation_accepts_multiple_files_and_sorts_groups_independently(client, monkeypatch):
    async def fake_convert_entries(entries, documents, job_dir, settings, progress_callback=None):
        pdf_paths = []
        for entry in entries:
            path = Path(job_dir) / "converted" / f"{entry['doc_id']}.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=595, height=842)
            with path.open("wb") as handle:
                writer.write(handle)
            entry["conversion_status"] = "Готов"
            entry["pages"] = 1
            entry["page_size"] = "210.0 × 297.0 мм · A4, книжная"
            entry["pdf_path"] = str(path.relative_to(job_dir))
            entry["pdf_analysis"] = {
                "page_count": 1,
                "pages": [{"width_mm": 210.0, "height_mm": 297.0, "rotate": 0}],
                "empty_pages": [],
            }
            pdf_paths.append(path)
        return entries, pdf_paths

    monkeypatch.setattr(main, "convert_entries", fake_convert_entries)
    upload = client.post(
        "/api/print/jobs",
        files=[
            ("files", ("Акт_ии1.docx", b"1", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("files", ("Акт_ии2.docx", b"2", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("files", ("Акт_ии10.docx", b"10", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("files", ("Акт_нн2.docx", b"n2", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("files", ("Акт_нн10.docx", b"n10", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ],
    )
    job_id = upload.json()["id"]

    def workbook_bytes(values):
        workbook = Workbook()
        sheet = workbook.active
        for row, value in enumerate(values, start=1):
            sheet.cell(row=row, column=1, value=value)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    response = client.post(
        f"/api/print/jobs/{job_id}/validate/excel",
        files=[
            ("files", ("first.xlsx", workbook_bytes(["ии10", "ии2", "ии1"]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
            ("files", ("second.xlsx", workbook_bytes(["нн10", "нн2"]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
        ],
        data={"stamping_json": '{"enabled": false}'},
    )

    assert response.status_code == 200
    validation = response.json()
    assert [item["name"] for item in validation["excel_files"]] == ["first.xlsx", "second.xlsx"]
    assert [group["column"] for group in validation["groups"]] == ["A", "A"]
    assert len({group["id"] for group in validation["groups"]}) == 2

    sort_response = client.post(
        f"/api/print/jobs/{job_id}/sort/excel",
        json={
            "groups": {
                validation["groups"][0]["id"]: ["ии1", "ии2", "ии10"],
                validation["groups"][1]["id"]: ["нн2", "нн10"],
            },
            "stamping": {"enabled": False},
        },
    )

    assert sort_response.status_code == 200
    sorted_groups = sort_response.json()["groups"]
    assert [entry["number"] for entry in sorted_groups[0]["validation"]["entries"]] == ["ии1", "ии2", "ии10"]
    assert [entry["number"] for entry in sorted_groups[1]["validation"]["entries"]] == ["нн2", "нн10"]

    build = client.post(f"/api/print/jobs/{job_id}/build?wait=true")

    assert build.status_code == 200
    assert [item["download_name"] for item in build.json()["result_pdfs"]] == [
        "1-(ии1-ии10).pdf",
        "2-(нн2-нн10).pdf",
    ]
    state = main.load_state(job_id)
    workbook = load_workbook(main.get_job_dir(job_id) / state["number_mapping_xlsx"], read_only=True, data_only=True)
    assert list(workbook["Сопоставление"].iter_rows(values_only=True))[1:] == [
        ("ии1", None),
        ("ии2", None),
        ("ии10", None),
        ("нн2", None),
        ("нн10", None),
    ]
    workbook.close()


def test_download_excel_part_pdf(client):
    upload = upload_docx(client)
    job_id = upload.json()["id"]
    job_dir = get_job_dir(job_id)
    part_path = job_dir / "result" / "pdf_parts" / "01_part.pdf"
    part_path.parent.mkdir(parents=True, exist_ok=True)
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    with part_path.open("wb") as handle:
        writer.write(handle)

    state = main.load_state(job_id)
    state["status"] = "ready"
    state["build"] = {
        "progress": "ready",
        "result_pdfs": [
            {
                "title": "Столбец A",
                "download_name": "01_part.pdf",
                "path": "result/pdf_parts/01_part.pdf",
                "page_count": 1,
                "size_bytes": part_path.stat().st_size,
            }
        ],
    }
    save_state(job_id, state)

    response = client.get(f"/api/print/jobs/{job_id}/download/part/1")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"


def test_conversion_error_is_reported(client, monkeypatch):
    async def fake_convert_entries(entries, documents, job_dir, settings, progress_callback=None):
        raise main.ConversionError("LibreOffice не создал PDF")

    monkeypatch.setattr(main, "convert_entries", fake_convert_entries)
    upload = upload_docx(client)
    job_id = upload.json()["id"]
    client.post(f"/api/print/jobs/{job_id}/validate", json={"sequence": "ее5968"})
    response = client.post(f"/api/print/jobs/{job_id}/build?wait=true")
    assert response.status_code == 400
    assert "LibreOffice" in response.json()["detail"]
