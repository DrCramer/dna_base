import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.print_service.services.report_service import write_number_mapping_xlsx


@pytest.mark.parametrize("sizes", [(3, 1), (100, 7), (101, 2)])
def test_mapping_blocks_match_actual_party_sizes_and_preserve_order(tmp_path, sizes):
    path = tmp_path / "mapping.xlsx"
    entries = []
    for part, size in enumerate(sizes, start=1):
        for position in range(1, size + 1):
            entries.append({
                "source_number_original": f"ии{size - position + 1}",
                "assigned_number": f"{len(entries) + 1:04d}-2026",
                "party_no": str(204 + part),
                "pdf_name": f"{part}.pdf",
                "pdf_position": position,
            })

    assert write_number_mapping_xlsx(entries, path) == sum(sizes)

    workbook = load_workbook(path)
    sheet = workbook["Сопоставление"]
    assert workbook.sheetnames == ["Сопоставление", "Детали"]
    assert sheet.freeze_panes == "A3"
    assert sheet.max_row == max(sizes) + 2
    assert sheet.max_column == 5
    for part, size in enumerate(sizes):
        column = part * 3 + 1
        selected = entries[sum(sizes[:part]):sum(sizes[:part + 1])]
        assert sheet.cell(1, column).value == f"Партия {205 + part} · документов: {size}"
        assert sheet.cell(2, column).value == "Исходный номер"
        assert sheet.cell(2, column + 1).value == "Присвоенный номер"
        assert list(sheet.iter_rows(min_row=3, max_row=size + 2, min_col=column, max_col=column + 1, values_only=True)) == [
            (entry["source_number_original"], entry["assigned_number"]) for entry in selected
        ]
    assert all(row[0] is None for row in sheet.iter_rows(min_col=3, max_col=3, values_only=True))
    assert sheet["D" + str(sizes[1] + 3)].value is None
    details = list(workbook["Детали"].iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in details] == [entry["source_number_original"] for entry in entries]
    assert [row[3:] for row in details] == [(entry["pdf_name"], entry["pdf_position"]) for entry in entries]
    workbook.close()


def test_duplicate_group_titles_do_not_merge_pdf_blocks_and_columns_can_exceed_z(tmp_path):
    path = tmp_path / "mapping.xlsx"
    entries = [{"number": f"ии{index}", "assigned_number": f"{index}-2026", "group": "Столбец A", "pdf_name": f"{index}.pdf"} for index in range(1, 13)]

    write_number_mapping_xlsx(entries, path)

    workbook = load_workbook(path)
    sheet = workbook["Сопоставление"]
    for index in range(12):
        column = index * 3 + 1
        assert sheet.cell(3, column).value == f"ии{index + 1}"
        assert sheet.column_dimensions[get_column_letter(column)].width == 28
    assert len(sheet.merged_cells.ranges) == 12
    workbook.close()


@pytest.mark.parametrize("entries", [[], [{"number": "ии1", "stamp_label": "0001-2026", "pdf_name": "one.pdf"}]])
def test_single_pdf_and_empty_mapping_keep_two_columns(tmp_path, entries):
    path = tmp_path / "mapping.xlsx"

    assert write_number_mapping_xlsx(entries, path) == len(entries)

    workbook = load_workbook(path)
    sheet = workbook["Сопоставление"]
    assert sheet.freeze_panes == "A2"
    assert list(sheet.values) == [("Исходный номер", "Присвоенный номер")] + [("ии1", "0001-2026")] * len(entries)
    assert not sheet.merged_cells.ranges
    workbook.close()
