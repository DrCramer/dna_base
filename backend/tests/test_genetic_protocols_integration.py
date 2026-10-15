import asyncio
from io import BytesIO
import os
import json

from alembic import command
from alembic.config import Config
import httpx
from openpyxl import Workbook, load_workbook
import pytest
import pytest_asyncio
from sqlalchemy import select, text
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import NullPool

from app.core.config import get_settings
from app.models import User, UserRole


@pytest.fixture(scope="module")
def isolated_url():
    database = os.environ.get("DNA_TEST_DATABASE_NAME")
    if not database or not database.startswith("dna_genetic_test_"):
        pytest.skip("Нужна отдельная БД DNA_TEST_DATABASE_NAME=dna_genetic_test_...")
    old_url = os.environ.get("DATABASE_URL")
    url = make_url(get_settings().database_url).set(database=database)
    os.environ["DATABASE_URL"] = url.render_as_string(hide_password=False)
    get_settings.cache_clear()
    cfg = Config("alembic.ini")
    command.upgrade(cfg, "0013_lab_protocols")

    async def legacy_seed():
        engine = create_async_engine(url)
        async with engine.begin() as connection:
            await connection.execute(text("INSERT INTO lab_protocols (series_key,protocol_no,protocol_date,name,snapshot_json) VALUES ('legacy',1,'2026-10-14','Legacy', CAST(:snapshot AS jsonb))"), {"snapshot": json.dumps({"schema_version": 1, "objects": [], "stages": []})})
        await engine.dispose()
    asyncio.run(legacy_seed())
    command.upgrade(cfg, "head")
    command.downgrade(cfg, "0013_lab_protocols")
    command.upgrade(cfg, "head")
    yield url
    if old_url is None:
        os.environ.pop("DATABASE_URL", None)
    else:
        os.environ["DATABASE_URL"] = old_url
    get_settings.cache_clear()


@pytest_asyncio.fixture
async def client(isolated_url):
    from app.main import app
    from app.api.deps import current_user, db_session, edit_user, admin_user
    engine = create_async_engine(isolated_url, poolclass=NullPool)
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    async with sessions() as session:
        user = (await session.execute(select(User).where(User.username == "genetic_test"))).scalar_one_or_none()
        if user is None:
            user = User(username="genetic_test", password_hash="test-only", role=UserRole.admin, is_active=True)
            session.add(user)
            await session.commit()
        async def database_session():
            async with sessions() as request_session:
                yield request_session
        async def auth_user():
            return user
        app.dependency_overrides.update({db_session: database_session, current_user: auth_user, edit_user: auth_user, admin_user: auth_user})
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as http:
            yield http, engine
    app.dependency_overrides.clear()
    await engine.dispose()


async def create_objects(http, year, names):
    response = await http.post("/api/protocols/genetic-examinations/numbers/apply", json={"examination_year": year, "items": names})
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.asyncio
async def test_migration_preserves_legacy_and_independent_numbering(client):
    http, engine = client
    old = (await http.get("/api/protocols?kind=chs&limit=200")).json()["items"]
    assert any(item["name"] == "Legacy" and item["kind"] == "chs" for item in old)
    assert (await http.get("/api/protocols/meta?kind=chs&protocol_date=2026-10-14")).json()["suggested_no"] == 2
    assert (await http.get("/api/protocols/meta?kind=genetic_examination&protocol_date=2026-10-14")).json()["suggested_no"] == 1


@pytest.mark.asyncio
async def test_identity_years_repeated_import_and_no_registry_rows(client):
    http, engine = client
    names = [{"raw_name": name} for name in ("251-1", "251-2", "251-80", "252-1")]
    current = await create_objects(http, 2026, names)
    previous = await create_objects(http, 2025, names)
    repeated = await create_objects(http, 2026, names)
    assert current["added"] == 4 and repeated["existing"] == 4 and repeated["added"] == 0
    assert current["object_ids"] == repeated["object_ids"]
    assert not set(current["object_ids"]) & set(previous["object_ids"])
    assert {obj["examination_no"] for obj in current["objects"]} == {"251", "252"}
    assert all("party_id" not in obj and "rcsme_reg_no" not in obj for obj in current["objects"])
    async with engine.connect() as connection:
        for table in ("objects", "parties", "stage_events", "rt_results", "realtime_details"):
            assert (await connection.execute(text(f"SELECT count(*) FROM {table}"))).scalar_one() == 0
    bad = await http.post("/api/protocols/preview", json={"kind": "genetic_examination", "examination_year": 2024, "protocol_date": "2026-10-14", "protocol_no": 1, "name": "wrong year", "object_ids": current["object_ids"]})
    assert bad.status_code == 400


@pytest.mark.asyncio
async def test_rt_multiple_files_targets_snapshot_revision_and_duplicate_guard(client):
    http, engine = client
    objects = await create_objects(http, 2026, [{"raw_name": "RT-251/1", "examination_no": "RT251"}, {"raw_name": "RT-251/2", "examination_no": "RT251"}])
    ids = objects["object_ids"]
    csv_one = b"Sample Name,Target Name,Quantity,Ct\nRT-251/1,Long,0.5,27\nRT-251/1,Small,0.7,26\nRT-251/1,Y,0.02,29\nUnmatched,Small,1,25\n"
    csv_two = b"Sample Name,Target Name,Quantity,Ct\nRT-251/2,Long,1.2,27\nRT-251/2,Small,0.8,26\nRT-251/2,Y,0.1,29\n"
    upload = await http.post("/api/protocols/genetic-examinations/rt/preview", data={"examination_year": "2026", "object_ids": ",".join(map(str, ids))}, files=[("files", ("one.csv", csv_one, "text/csv")), ("files", ("two.csv", csv_two, "text/csv"))])
    assert upload.status_code == 200, upload.text
    files = upload.json()["files"]
    assert len(files) == 2
    assert files[0]["samples"][0]["small_quantity"] == 0.7
    assert files[0]["samples"][1]["status"] == "unmatched"
    commit = {"examination_year": 2026, "object_ids": ids, "upload_ids": [item["upload_id"] for item in files]}
    response = await http.post("/api/protocols/genetic-examinations/rt/commit", json=commit)
    assert response.status_code == 200, response.text
    assert response.json()["matched"] == 2
    duplicate = await http.post("/api/protocols/genetic-examinations/rt/commit", json=commit)
    assert duplicate.status_code == 409
    confirmed = await http.post("/api/protocols/genetic-examinations/rt/commit", json={**commit, "confirm_duplicate": True})
    assert confirmed.status_code == 200 and confirmed.json()["matched"] == 2
    payload = {"kind": "genetic_examination", "examination_year": 2026, "protocol_date": "2026-10-14", "protocol_no": 1, "name": "GeneticRT", "object_ids": ids, "dilution": {"show_all_targets": True, "target_overrides": {f"genetic_examination:{ids[0]}": "long"}}}
    saved = await http.post("/api/protocols", json=payload)
    assert saved.status_code == 200, saved.text
    protocol = saved.json()
    snapshot = protocol["snapshot"]
    assert snapshot["schema_version"] == 2 and snapshot["kind"] == "genetic_examination"
    assert snapshot["dilutions"][0]["source_concentration"] == 0.5
    assert snapshot["dilutions"][1]["source_concentration"] == 0.8
    protocol_id = protocol["id"]
    assert (await http.post(f"/api/protocols/{protocol_id}/finalize")).status_code == 200
    revision = (await http.post(f"/api/protocols/{protocol_id}/revision")).json()
    assert revision["kind"] == "genetic_examination" and revision["revision_no"] == 2
    assert revision["snapshot"] == snapshot
    workbook = load_workbook(BytesIO((await http.get(f"/api/protocols/{protocol_id}/excel")).content), data_only=True)
    assert "Мишень" in [cell.value for cell in workbook["Dilution"][1]]
    assert (await http.get(f"/api/protocols/{protocol_id}")).json()["snapshot"] == snapshot
    duplicate_protocol = (await http.post(f"/api/protocols/{protocol_id}/duplicate", json={"copy_objects": True})).json()
    assert duplicate_protocol["kind"] == "genetic_examination"
    assert duplicate_protocol["snapshot"]["objects"] == snapshot["objects"]
    async with engine.connect() as connection:
        rows = (await connection.execute(text("SELECT object_id,genetic_examination_object_id,party_id FROM lab_protocol_objects WHERE protocol_id=:id"), {"id": protocol_id})).all()
        assert len(rows) == 2 and all(row[0] is None and row[1] in ids and row[2] is None for row in rows)
        wells = (await connection.execute(text("SELECT object_id,genetic_examination_object_id FROM lab_protocol_wells WHERE protocol_id=:id AND kind='sample'"), {"id": protocol_id})).all()
        assert all(row[0] is None and row[1] in ids for row in wells)
        assert (await connection.execute(text("SELECT count(*) FROM stage_events"))).scalar_one() == 0


@pytest.mark.asyncio
async def test_excel_adapter_preserves_coordinates_and_unknown_performers(client):
    http, engine = client
    workbook = Workbook()
    source = workbook.active
    source.title = "Единая_плашка_A4"
    source["A1"] = "Единый протокол"
    source["A2"], source["B2"] = "Дата", "14.10.2026"
    source["D2"], source["E2"] = "№", "11"
    source["F2"], source["G2"] = "Название", "Imported"
    source["B4"], source["B5"], source["B6"] = "14.10.2026", "Неизвестный набор", "Неизвестный исполнитель"
    source["A10"], source["B10"], source["C10"] = "A", "EX-1", "EX-2"
    pcr = workbook.create_sheet("PCR")
    pcr["A18"], pcr["B18"] = "A", "L1"
    pcr["A19"], pcr["B19"] = "B", "EX-2"
    pcr["A25"], pcr["M25"] = "H", "EX-1"
    content = BytesIO(); workbook.save(content)
    response = await http.post("/api/protocols/genetic-examinations/excel/preview", data={"examination_year": "2026"}, files={"file": ("protocol.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview["unresolved"] == 2
    payload = {"examination_year": 2026, "upload_id": preview["upload_id"], "items": [{"raw_name": item["raw_name"], "examination_no": "EX"} for item in preview["items"]]}
    applied = await http.post("/api/protocols/genetic-examinations/excel/apply", json=payload)
    assert applied.status_code == 200, applied.text
    data = applied.json()
    assert data["warnings"] and data["stages"][0]["performers"][0]["display_name"] == "Неизвестный исполнитель"
    request = {"kind": "genetic_examination", "examination_year": 2026, "protocol_date": "2026-10-14", "protocol_no": 11, "name": "Imported", "object_ids": data["object_ids"], "imported_wells": data["imported_wells"], "stages": data["stages"]}
    preview_protocol = await http.post("/api/protocols/preview", json=request)
    assert preview_protocol.status_code == 200, preview_protocol.text
    wells = preview_protocol.json()["layouts"]["pcr"]["plates"][0]["wells"]
    assert next(well for well in wells if well["display_name"] == "EX-1")["well"] == "H12"
    assert next(well for well in wells if well["display_name"] == "EX-2")["well"] == "B1"
    assert preview_protocol.json()["stages"][0]["performers"][0]["display_name"] == "Неизвестный исполнитель"


@pytest.mark.asyncio
async def test_excel_duplicate_physical_objects_manual_well_mapping_and_selected_fields(client):
    http, engine = client
    workbook = Workbook()
    source = workbook.active
    source.title = "Единая_плашка_A4"
    source["B4"], source["B5"] = "14.10.2026", "ImportedKit"
    source["A10"], source["B10"] = "A", "SAME"
    source["A11"], source["B11"] = "B", "SAME"
    pcr = workbook.create_sheet("PCR")
    pcr["A18"], pcr["B18"] = "A", "SAME"
    pcr["A19"], pcr["B19"] = "B", "SAME"
    content = BytesIO(); workbook.save(content)
    response = await http.post("/api/protocols/genetic-examinations/excel/preview", data={"examination_year": "2026"}, files={"file": ("duplicate.xlsx", content.getvalue())})
    assert response.status_code == 200, response.text
    preview = response.json()
    assert len(preview["items"]) == 2 and preview["unresolved"] == 2
    payload = {"examination_year": 2026, "upload_id": preview["upload_id"], "selected_stage_fields": ["kit"], "items": [{"source_key": item["source_key"], "raw_name": f"Corrected-{index}", "examination_no": "Corrected"} for index, item in enumerate(preview["items"], 1)]}
    assert (await http.post("/api/protocols/genetic-examinations/excel/apply", json=payload)).status_code == 400
    payload["well_mappings"] = {"pcr:1:A1": "source:1:B1", "pcr:1:B1": "source:1:A1"}
    response = await http.post("/api/protocols/genetic-examinations/excel/apply", json=payload)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["added"] == 2  # The failed transaction did not create either object.
    assert result["stages"][0]["kit_name"] == "ImportedKit"
    assert "work_date" not in result["stages"][0] and "performers" not in result["stages"][0]
    assert next(cell for cell in result["imported_wells"] if cell["layout_key"] == "pcr" and cell["well"] == "A1")["source_id"] == result["object_ids"][1]


@pytest.mark.asyncio
async def test_rt_atomic_batch_latest_attempt_and_frozen_snapshot(client):
    http, engine = client
    objects = await create_objects(http, 2027, [{"raw_name": "AT-1", "examination_no": "AT"}, {"raw_name": "AT-2", "examination_no": "AT"}])
    ids = objects["object_ids"]
    async def upload(content, filename="attempt.csv"):
        response = await http.post("/api/protocols/genetic-examinations/rt/preview", data={"examination_year": "2027", "object_ids": ",".join(map(str, ids))}, files=[("files", (filename, content, "text/csv"))])
        assert response.status_code == 200, response.text
        return response.json()["files"][0]["upload_id"]
    first = await upload(b"Sample Name,Target Name,Quantity,Ct\nAT-1,Long,0.5,27\nAT-1,Small,0.7,26\nAT-2,Small,0.8,26\n")
    commit = {"examination_year": 2027, "object_ids": ids, "upload_ids": [first]}
    assert (await http.post("/api/protocols/genetic-examinations/rt/commit", json=commit)).status_code == 200
    payload = {"kind": "genetic_examination", "examination_year": 2027, "protocol_date": "2026-10-14", "protocol_no": 30, "name": "FrozenRT", "object_ids": ids}
    saved = (await http.post("/api/protocols", json=payload)).json()
    valid = await upload(b"Sample Name,Target Name,Quantity,Ct\nAT-1,Long,0.25,24\n", "new.csv")
    invalid = await upload(b"Sample Name,Target Name,Quantity,Ct\nUnknown,Small,1,24\n", "unmatched.csv")
    async with engine.connect() as connection:
        before = (await connection.execute(text("SELECT count(*) FROM genetic_examination_rt_runs"))).scalar_one()
    failure = await http.post("/api/protocols/genetic-examinations/rt/commit", json={**commit, "upload_ids": [valid, invalid], "mode": "replace"})
    assert failure.status_code == 400
    async with engine.connect() as connection:
        assert (await connection.execute(text("SELECT count(*) FROM genetic_examination_rt_runs"))).scalar_one() == before
    response = await http.post("/api/protocols/genetic-examinations/rt/commit", json={**commit, "upload_ids": [valid], "mode": "replace"})
    assert response.status_code == 200, response.text
    rows = response.json()["objects"]
    assert rows[0]["rt"]["long_quantity"] == 0.25 and rows[0]["rt"]["small_quantity"] is None
    assert rows[1]["rt"]["small_quantity"] == 0.8
    preview = (await http.post("/api/protocols/preview", json=payload)).json()
    assert preview["dilutions"][0]["available"] is False
    assert (await http.get(f"/api/protocols/{saved['id']}")).json()["snapshot"] == saved["snapshot"]
    async with engine.connect() as connection:
        assert (await connection.execute(text("SELECT count(*) FROM genetic_examination_rt_results WHERE genetic_examination_object_id=:id AND is_current=false"), {"id": ids[0]})).scalar_one() == 2


@pytest.mark.asyncio
async def test_rt_manual_ambiguous_match_is_year_scoped_and_upload_owner_checked(client):
    from app.api.deps import edit_user
    from app.main import app
    http, engine = client
    objects = await create_objects(http, 2028, [{"raw_name": "SAMPLE/1", "examination_no": "First"}, {"raw_name": "SAMPLE/1", "examination_no": "Second"}])
    other = await create_objects(http, 2029, [{"raw_name": "SAMPLE/1", "examination_no": "First"}])
    ids = objects["object_ids"]
    response = await http.post("/api/protocols/genetic-examinations/rt/preview", data={"examination_year": "2028", "object_ids": ",".join(map(str, ids))}, files=[("files", ("ambiguous.csv", b"Sample Name,Target Name,Quantity,Ct\nSAMPLE/1,Small,0.5,27\n", "text/csv"))])
    file = response.json()["files"][0]
    assert file["samples"][0]["status"] == "ambiguous"
    payload = {"examination_year": 2028, "object_ids": ids, "upload_ids": [file["upload_id"]], "mappings": [{"upload_id": file["upload_id"], "sample_key": "sample/1", "object_id": other["object_ids"][0]}]}
    assert (await http.post("/api/protocols/genetic-examinations/rt/commit", json=payload)).status_code == 400
    payload["mappings"][0]["object_id"] = ids[1]
    original_auth = app.dependency_overrides[edit_user]
    async def other_user():
        return User(id=999999, username="not-owner", role=UserRole.user)
    app.dependency_overrides[edit_user] = other_user
    assert (await http.post("/api/protocols/genetic-examinations/rt/commit", json=payload)).status_code == 404
    app.dependency_overrides[edit_user] = original_auth
    applied = await http.post("/api/protocols/genetic-examinations/rt/commit", json=payload)
    assert applied.status_code == 200, applied.text
    assert applied.json()["objects"][0]["rt"] is None
    assert applied.json()["objects"][1]["rt"]["small_quantity"] == 0.5


@pytest.mark.asyncio
async def test_duplicate_settings_retains_year_but_clears_selection(client):
    http, engine = client
    objects = await create_objects(http, 2024, [{"raw_name": "DUP-1", "examination_no": "DUP"}])
    payload = {"kind": "genetic_examination", "examination_year": 2024, "protocol_date": "2026-10-14", "protocol_no": 40, "name": "CopySettings", "object_ids": objects["object_ids"], "dilution": {"target_overrides": {f"genetic_examination:{objects['object_ids'][0]}": "y"}}}
    saved = (await http.post("/api/protocols", json=payload)).json()
    copy = (await http.post(f"/api/protocols/{saved['id']}/duplicate", json={"copy_objects": False})).json()
    assert copy["snapshot"]["examination_year"] == 2024
    assert copy["snapshot"]["objects"] == [] and copy["snapshot"]["imported_wells"] == []
    assert copy["snapshot"]["dilution_settings"]["target_overrides"] == {}
    assert copy["snapshot"]["selection"] is None and copy["snapshot"]["dilution_optimization"] is None
    assert (await http.get(f"/api/protocols/{copy['id']}/excel")).status_code == 200
    assert copy["id"] in {item["id"] for item in (await http.get("/api/protocols?kind=genetic_examination&year=2024")).json()["items"]}
    assert saved["id"] in {item["id"] for item in (await http.get("/api/protocols?kind=genetic_examination&examination_no=DUP&object_no=DUP-1")).json()["items"]}


@pytest.mark.asyncio
async def test_genetic_rounding_applies_to_preview_calculation_snapshot_and_excel_not_raw_measurement(client):
    http, engine = client
    created = await create_objects(http, 2031, [{"raw_name": "PRECISION-1", "examination_no": "PRECISION"}])
    ids = created["object_ids"]
    response = await http.post("/api/protocols/genetic-examinations/rt/preview", data={"examination_year": "2031", "object_ids": str(ids[0])}, files=[("files", ("precision.csv", b"Sample Name,Target Name,Quantity,Ct\nPRECISION-1,Long,0.81234567,27\nPRECISION-1,Small,0.63825689,26\nPRECISION-1,Y,0.01429445,29\n", "text/csv"))])
    assert response.status_code == 200, response.text
    file = response.json()["files"][0]
    assert file["samples"][0]["small_quantity"] == 0.6383
    committed = await http.post("/api/protocols/genetic-examinations/rt/commit", json={"examination_year": 2031, "object_ids": ids, "upload_ids": [file["upload_id"]]})
    assert committed.status_code == 200, committed.text
    assert committed.json()["objects"][0]["rt"]["small_quantity"] == 0.6383
    payload = {"kind": "genetic_examination", "examination_year": 2031, "protocol_date": "2026-10-15", "protocol_no": 80, "name": "Precision", "object_ids": ids, "dilution": {"show_all_targets": True}}
    saved = await http.post("/api/protocols", json=payload)
    assert saved.status_code == 200, saved.text
    protocol = saved.json()
    row = protocol["snapshot"]["dilutions"][0]
    assert row["source_concentration"] == 0.6383
    assert row["rt_targets"] == {"long": 0.8123, "small": 0.6383, "y": 0.0143}
    workbook = load_workbook(BytesIO((await http.get(f"/api/protocols/{protocol['id']}/excel")).content), data_only=True)
    assert workbook["Dilution"]["D2"].value == 0.6383
    async with engine.connect() as connection:
        assert (await connection.execute(text("SELECT quantity FROM genetic_examination_rt_results WHERE genetic_examination_object_id=:id AND target_key='small'"), {"id": ids[0]})).scalar_one() == 0.63825689


@pytest.mark.asyncio
async def test_viewer_can_read_but_cannot_import_or_change_protocols(client):
    from app.api.deps import current_user, edit_user, admin_user
    from app.main import app
    http, engine = client
    objects = await create_objects(http, 2030, [{"raw_name": "VIEW-1", "examination_no": "VIEW"}])
    saved = await http.post("/api/protocols", json={"kind": "genetic_examination", "examination_year": 2030, "protocol_date": "2026-10-14", "protocol_no": 70, "name": "ViewerRead", "object_ids": objects["object_ids"]})
    assert saved.status_code == 200, saved.text
    protocol_id = saved.json()["id"]
    old_overrides = dict(app.dependency_overrides)
    async def viewer():
        return User(id=999998, username="viewer-test", role=UserRole.viewer)
    app.dependency_overrides[current_user] = viewer
    app.dependency_overrides.pop(edit_user)
    app.dependency_overrides.pop(admin_user)
    try:
        assert (await http.get("/api/protocols?kind=genetic_examination")).status_code == 200
        assert (await http.get("/api/protocols/1")).status_code == 200
        assert (await http.get(f"/api/protocols/{protocol_id}/excel")).status_code == 200
        assert (await http.post("/api/protocols/genetic-examinations/numbers/apply", json={"examination_year": 2026, "items": [{"raw_name": "999-1"}]})).status_code == 403
        assert (await http.post("/api/protocols/1/finalize")).status_code == 403
        assert (await http.post("/api/protocols/1/archive")).status_code == 403
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(old_overrides)
