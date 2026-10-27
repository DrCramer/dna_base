import copy
import math
import time
from io import BytesIO
from unittest.mock import AsyncMock

import pytest
from pydantic import ValidationError
from openpyxl import load_workbook

from app.schemas.protocols import ProtocolDilutionSettings
from app.schemas.protocols import ProtocolPreviewRequest, ProtocolPreviewOut
from app.services import protocols
from app.services.protocol_excel import build_protocol_workbook
from app.services.protocol_calculations import calculate_dilution
from app.services.protocol_dilution_optimizer import (
    calculate_actual_dilution_concentration,
    generate_allowed_water_candidates,
    optimize_dilution_water_batches,
)


def settings(**patch):
    return ProtocolDilutionSettings(optimize_water_volumes=True, **patch)


def row(source, index=1, **kwargs):
    return {"object_id": index, "display_name": f"ии{index}", "plate_index": 1 + index // 88, "well": "B2", **calculate_dilution(source, **kwargs)}


@pytest.mark.parametrize("target", [0.2, 0.5, 1, 1.5])
def test_optimizer_checks_the_changed_target_and_preserves_minimum_and_source_limits(target):
    rows = [row(source, index, target_concentration=target, minimum_final_volume_enabled=True) for index, source in enumerate([10, 1000, 0.4], 1)]
    policy = settings(target_concentration=target, minimum_final_volume_enabled=True)
    optimized, summary = optimize_dilution_water_batches(rows, policy)
    assert summary["enabled"]
    for item in optimized:
        assert item["target_concentration"] == target
        if not item["steps"]:
            assert item["source_concentration"] <= target
            continue
        metadata = item["optimization"]
        actual = calculate_actual_dilution_concentration(item["source_concentration"], item["steps"])
        assert metadata["target_concentration"] == target
        assert metadata["actual_final_concentration"] == pytest.approx(actual)
        if not metadata["needs_review"]:
            assert abs(100 * (actual / target - 1)) <= 5 + 1e-9
            assert item["steps"][0]["dna_volume"] <= 50
            assert all(step["dna_volume"] + step["water_volume"] >= 15 for step in item["steps"])


@pytest.mark.parametrize("factor,water,accepted,expected_error", [
    (82.72, 250, True, -1.91), (43.96, 130, True, -0.84),
    (22.47, 70, False, -7.66), (37.38, 100, False, 8.87),
    (38.44, 100, False, 11.96), (9.231, 30, False, -16.08),
])
def test_photo_examples_are_checked_by_actual_concentration(factor, water, accepted, expected_error):
    actual = calculate_actual_dilution_concentration(factor * 0.1, [{"dna_volume": 3, "water_volume": water}])
    error = 100 * (actual / 0.1 - 1)
    assert error == pytest.approx(expected_error, abs=0.02)
    candidates = generate_allowed_water_candidates(factor * 0.1, 0.1, 3, baseline_water_volume=math.ceil((factor - 1) * 3), settings=settings())
    assert (float(water) in candidates) is accepted


def test_both_upward_and_downward_candidates_can_be_accepted():
    candidates = generate_allowed_water_candidates(8.272, 0.1, 3, baseline_water_volume=246, settings=settings())
    assert 240 in candidates
    assert 250 in candidates


@pytest.mark.parametrize("baseline_water,step,rounded_water", [(242, 10, 240), (247, 10, 250), (242, 5, 240), (247, 5, 245), (248, 5, 250)])
def test_single_sample_prefers_nearest_convenient_volume_in_either_direction(baseline_water, step, rounded_water):
    source = row((baseline_water + 3) / 30)
    result, _ = optimize_dilution_water_batches([source], settings(preferred_water_step_ul=step))
    item = result[0]
    assert item["baseline_steps"][0]["water_volume"] == baseline_water
    assert item["steps"][0]["water_volume"] == rounded_water
    assert item["steps"][0]["dna_volume"] == source["steps"][0]["dna_volume"]
    assert item["optimization"]["status"] == "optimized"
    assert abs(item["optimization"]["relative_error_percent"]) <= 5


def test_batch_rounds_both_up_and_down_without_forcing_the_same_direction():
    rows = [row((water + 3) / 30, index, threshold=1000) for index, water in enumerate([242, 247, 483, 487], start=1)]
    result, _ = optimize_dilution_water_batches(rows, settings())
    water = [item["steps"][0]["water_volume"] for item in result]
    assert water[:2] == [240, 240]
    assert water[2:] == [490, 490]
    assert any(after < before["steps"][0]["water_volume"] for before, after in zip(rows, water))
    assert any(after > before["steps"][0]["water_volume"] for before, after in zip(rows, water))
    assert all(value % 10 == 0 for value in water)
    assert optimize_dilution_water_batches(rows, settings())[0] == result


@pytest.mark.parametrize("concentration", [117, 1000])
def test_two_step_rounding_can_use_opposite_directions_with_valid_final_concentration(concentration):
    source = row(concentration)
    result, _ = optimize_dilution_water_batches([source], settings())
    item = result[0]
    differences = [after["water_volume"] - before["water_volume"] for before, after in zip(source["steps"], item["steps"])]
    assert differences[0] * differences[1] < 0
    assert all(step["water_volume"] % 10 == 0 for step in item["steps"])
    assert [step["dna_volume"] for step in item["steps"]] == [step["dna_volume"] for step in source["steps"]]
    assert abs(item["optimization"]["relative_error_percent"]) <= 5
    assert item["optimization"]["actual_final_concentration"] == pytest.approx(calculate_actual_dilution_concentration(concentration, item["steps"]))


def test_both_convenient_directions_survive_dense_common_volume_candidates():
    candidates = generate_allowed_water_candidates(8.272, 0.1, 3, baseline_water_volume=246,
        settings=settings(pipetting={"water_increment_ul": 0.1}), common_volumes=[index / 10 for index in range(2300, 2601)])
    assert 240 in candidates
    assert 250 in candidates
    assert 246 in candidates
    assert len(candidates) <= 32


@pytest.mark.parametrize("limits,expected", [({"max_under_target_pct": 0}, 240), ({"max_over_target_pct": 0}, 250)])
def test_convenient_rounding_respects_one_sided_concentration_limits(limits, expected):
    result, _ = optimize_dilution_water_batches([row((247 + 3) / 30)], settings(**limits))
    assert result[0]["steps"][0]["water_volume"] == expected


def test_convenient_rounding_falls_back_when_both_neighbours_are_outside_tolerance():
    source = row((64 + 3) / 30)
    result, _ = optimize_dilution_water_batches([source], settings())
    assert result[0]["steps"][0]["water_volume"] == source["steps"][0]["water_volume"]
    assert result[0]["steps"][0]["dna_volume"] == source["steps"][0]["dna_volume"]
    assert result[0]["optimization"]["status"] == "unchanged"


def test_minimum_error_priority_keeps_accurate_nonrounded_baseline():
    source = row((247 + 3) / 30)
    result, _ = optimize_dilution_water_batches([source], settings(optimization_priority="minimum_error"))
    assert result[0]["steps"][0]["water_volume"] == source["steps"][0]["water_volume"]
    assert result[0]["steps"][0]["dna_volume"] == source["steps"][0]["dna_volume"]
    assert result[0]["optimization"]["relative_error_percent"] == pytest.approx(0)


def test_preferred_water_step_affects_candidates_and_can_be_extended():
    arguments = dict(source_concentration=8.272, target_concentration=0.1, dna_volume=3, baseline_water_volume=246)
    coarse = generate_allowed_water_candidates(**arguments, settings=settings(preferred_water_step_ul=10))
    fine = generate_allowed_water_candidates(**arguments, settings=settings(preferred_water_step_ul=5))
    assert 245 not in coarse
    assert 245 in fine
    assert ProtocolDilutionSettings(preferred_water_step_ul=2).preferred_water_step_ul == 2


def test_batch_groups_reduce_settings_without_changing_dna_or_identity():
    rows = [row((water + 3) / 30, index) for index, water in enumerate(range(63, 68), start=1)]
    original = copy.deepcopy(rows)

    result, summary = optimize_dilution_water_batches(rows, settings())

    assert rows == original
    assert summary["unique_water_volumes_after"] < summary["unique_water_volumes_before"]
    assert summary["blocked_object_count"] == 0
    assert {item["steps"][0]["water_volume"] for item in result} == {65}
    for before, after in zip(rows, result):
        assert {key: after[key] for key in ("object_id", "plate_index", "well")} == {key: before[key] for key in ("object_id", "plate_index", "well")}
        assert after["steps"][0]["dna_volume"] == before["steps"][0]["dna_volume"]
        assert after["baseline_steps"] == before["steps"]
        assert abs(after["optimization"]["relative_error_percent"]) <= 5 + 1e-9
        assert after["optimization"]["actual_final_concentration"] == calculate_actual_dilution_concentration(before["source_concentration"], after["steps"])
    assert optimize_dilution_water_batches(rows, settings()) == (result, summary)


def test_invalid_baseline_is_flagged_instead_of_reported_ready():
    result, summary = optimize_dilution_water_batches([row(0.108)], settings())
    item = result[0]
    assert item["steps"] == item["baseline_steps"]
    assert item["optimization"]["status"] == "blocked"
    assert item["optimization"]["needs_review"] is True
    assert item["optimization"]["baseline_out_of_tolerance"] is True
    assert item["optimization"]["reason_code"] == "baseline_out_of_tolerance"
    assert summary["blocked_object_count"] == 1


def test_baseline_can_be_repaired_using_fixed_dna_and_measurable_water():
    result, _ = optimize_dilution_water_batches([row(0.134)], settings())
    item = result[0]
    assert item["baseline_steps"][0]["water_volume"] == 2
    assert item["steps"][0]["water_volume"] == 1
    assert item["steps"][0]["dna_volume"] == 3
    assert item["optimization"]["status"] == "optimized"
    assert item["optimization"]["baseline_out_of_tolerance"] is True
    assert abs(item["optimization"]["relative_error_percent"]) < 5


def test_missing_small_and_unnecessary_dilution_are_excluded():
    rows = [row(None), row(0.05, 2)]
    result, summary = optimize_dilution_water_batches(rows, settings())
    assert result == rows
    assert summary["excluded_object_count"] == 2
    assert summary["unique_water_volumes_after"] == 0


def test_minimum_volume_is_only_enforced_when_enabled():
    source = row(0.2)
    result, _ = optimize_dilution_water_batches([source], settings())
    assert result[0]["steps"][0]["dna_volume"] == 3
    assert sum(result[0]["steps"][0][key] for key in ("dna_volume", "water_volume")) < 15
    scaled = row(0.2, minimum_final_volume_enabled=True)
    result, _ = optimize_dilution_water_batches([scaled], settings(minimum_final_volume_enabled=True))
    assert result[0]["steps"][0]["dna_volume"] == 8
    assert result[0]["steps"][0]["water_volume"] >= 7
    assert sum(result[0]["steps"][0][key] for key in ("dna_volume", "water_volume")) >= 15
    impossible = row(1.01, target_concentration=1, minimum_final_volume_enabled=True)
    result, summary = optimize_dilution_water_batches([impossible], settings(minimum_final_volume_enabled=True))
    assert result[0]["steps"] == impossible["steps"]
    assert summary["blocked_object_count"] == 1


def test_source_limit_is_not_applied_to_second_stage_product():
    source = row(1000, dilution_one_volume=60)
    result, _ = optimize_dilution_water_batches([source], settings())
    item = result[0]
    assert item["optimization"]["status"] != "blocked"
    assert item["steps"][0]["dna_volume"] <= 50
    assert item["steps"][1]["dna_volume"] == 60
    assert item["optimization"]["first_dilution_remaining_volume"] >= 0


def test_two_steps_check_final_error_not_each_nominal_step_separately():
    source = row(1000)
    result, _ = optimize_dilution_water_batches([source], settings())
    item = result[0]
    actual = 1000
    for step in item["steps"]:
        actual *= step["dna_volume"] / (step["dna_volume"] + step["water_volume"])
    assert item["optimization"]["actual_final_concentration"] == pytest.approx(actual)
    assert abs(100 * (actual / 0.1 - 1)) <= 5 + 1e-9
    bad = row(1000)
    bad["steps"][0]["water_volume"] = 310
    bad["steps"][1]["water_volume"] = 1030
    # Оба шага отдельно отклоняются менее чем на 5%, но вместе превышают допуск.
    assert abs(100 * (100 / ((3 + 310) / 3) - 1)) < 5
    assert abs(100 * (100 / ((10 + 1030) / 10) - 1)) < 5
    result, _ = optimize_dilution_water_batches([bad], settings())
    item = result[0]
    assert item["optimization"]["baseline_out_of_tolerance"] is True
    assert item["optimization"]["needs_review"] or abs(item["optimization"]["relative_error_percent"]) <= 5 + 1e-9


def test_pipetting_limits_and_reserve_do_not_trigger_hidden_dna_changes():
    source = row(1000)
    result, _ = optimize_dilution_water_batches([source], settings(pipetting={"water_max_ul": 100}))
    assert result[0]["optimization"]["status"] == "blocked"
    assert result[0]["steps"] == source["steps"]
    scaled = row(0.21, minimum_final_volume_enabled=True)
    result, _ = optimize_dilution_water_batches([scaled], settings(minimum_final_volume_enabled=True, pipetting={"dna_increment_ul": 1}))
    assert result[0]["optimization"]["status"] == "blocked"
    result, _ = optimize_dilution_water_batches([source], settings(pipetting={"first_dilution_reserve_ul": 200}))
    assert result[0]["optimization"]["first_dilution_remaining_volume"] >= 200


def test_minimum_error_priority_and_asymmetric_limits():
    rows = [row(8.272)]
    result, _ = optimize_dilution_water_batches(rows, settings(optimization_priority="minimum_error", max_over_target_pct=0))
    assert result[0]["optimization"]["relative_error_percent"] <= 1e-9
    result, _ = optimize_dilution_water_batches(rows, settings(max_under_target_pct=0))
    assert result[0]["optimization"]["relative_error_percent"] >= -1e-9


@pytest.mark.parametrize("count", [300, 2000])
def test_batch_optimizer_has_bounded_runtime_and_checks_every_row(count):
    rows = [row(0.11 + (index % 600) * 0.0763, index + 1) for index in range(count)]
    start = time.perf_counter()

    result, summary = optimize_dilution_water_batches(rows, settings())

    assert time.perf_counter() - start < 15
    assert len(result) == count
    assert summary["optimized_object_count"] + summary["unchanged_object_count"] + summary["blocked_object_count"] == count
    for item in result:
        if not item["optimization"]["needs_review"]:
            assert abs(item["optimization"]["relative_error_percent"]) <= 5 + 1e-9


@pytest.mark.parametrize("patch", [{"allowed_concentration_error_percent": 100}, {"allowed_concentration_error_percent": -1}, {"pipetting": {"water_min_ul": 20, "water_max_ul": 10}}, {"pipetting": {"water_increment_ul": 0}}, {"source_dna_volume": float('inf')}])
def test_invalid_policy_is_rejected(patch):
    with pytest.raises(ValidationError):
        settings(**patch)


@pytest.mark.asyncio
async def test_preview_snapshot_and_excel_use_same_volumes_without_changing_layout(monkeypatch):
    objects = [{"id": index, "rcsme_reg_no": f"ии{index}", "rt": {"concentration": (63 + index + 3) / 30}} for index in range(1, 6)]
    objects += [{"id": 6, "rcsme_reg_no": "ии6", "rt": {"concentration": 0.2}}]
    monkeypatch.setattr(protocols, "protocol_object_snapshots", AsyncMock(return_value=objects))
    stages = [{"stage_type": stage, "enabled": True, "sequencer_name": None} for stage in protocols.STAGE_ORDER]
    monkeypatch.setattr(protocols, "_stage_snapshots", AsyncMock(return_value=(stages, {})))
    payload = ProtocolPreviewRequest(protocol_date="2026-10-12", protocol_no=1, name="test", object_ids=list(range(1, 7)), dilution={"minimum_final_volume_enabled": True})
    baseline = await protocols.build_protocol_snapshot(None, payload)
    assert baseline["dilution_optimization"] is None
    assert "baseline_steps" not in baseline["dilutions"][0]
    payload.dilution.optimize_water_volumes = True

    preview = ProtocolPreviewOut(**await protocols.build_protocol_snapshot(None, payload)).model_dump(mode="json")

    assert preview["layouts"] == baseline["layouts"]
    assert preview["snapshot"]["dilutions"] == preview["dilutions"]
    assert preview["snapshot"]["dilution_optimization"] == preview["dilution_optimization"]
    assert preview["snapshot"]["dilution_settings"]["allowed_concentration_error_percent"] == 5
    changes = [after["steps"][0]["water_volume"] - before["steps"][0]["water_volume"] for before, after in zip(baseline["dilutions"], preview["dilutions"]) if before["steps"]]
    assert any(change > 0 for change in changes)
    assert any(change < 0 for change in changes)
    workbook = load_workbook(BytesIO(build_protocol_workbook(preview["snapshot"])), data_only=True)
    sheet = workbook["Dilution"]
    for excel_row, item in enumerate(preview["dilutions"], start=2):
        assert sheet.cell(excel_row, 9).value == item["steps"][0]["water_volume"]
        assert sheet.cell(excel_row, 16).value == item["baseline_steps"][0]["water_volume"]
        assert sheet.cell(excel_row, 13).value == pytest.approx(item["optimization"]["actual_final_concentration"])
        assert sheet.cell(excel_row, 14).value == pytest.approx(item["optimization"]["relative_error_percent"])
    assert "Округление" in workbook.sheetnames
    assert "Группы H₂O" in workbook.sheetnames
    workbook.close()
