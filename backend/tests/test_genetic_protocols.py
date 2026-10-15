import copy
from types import SimpleNamespace

import pytest
from fastapi import HTTPException

from app.parsers.genetic_identity import genetic_key, genetic_number_parts
from app.parsers.rt import aggregate_rt_rows, round_rt_quantity
from app.parsers.work_protocol import parse_work_protocol_preview
from app.schemas.genetic_protocols import GeneticExcelApply, GeneticNumbersRequest
from app.services.genetic_excel import excel_cells, excel_source_items, resolve_excel_items, selected_stage_patch
from app.schemas.protocols import ProtocolImportedWell, ProtocolPreviewRequest
from app.services.genetic_protocols import match_genetic_sample, number_rows
from app.services.protocol_calculations import calculate_dilution
from app.services.protocol_import_layout import apply_imported_layouts
from app.services.protocol_plate import build_protocol_layouts
from app.services.protocols import build_protocol_dilutions, protocol_object_count


@pytest.mark.parametrize("raw, exam, index", [("251-1", "251", 1), ("251-80", "251", 80), ("251-2026", None, None), ("АБ-125/1", None, None), ("A-12-3", None, None)])
def test_examination_number_is_inferred_only_when_unambiguous(raw, exam, index):
    assert genetic_number_parts(raw) == (exam, index)


def test_explicit_examination_context_and_safe_identity():
    assert genetic_number_parts("251-2026", "251") == ("251", 2026)
    assert genetic_key("  АБ–125/1 ") == "аб-125/1"
    assert genetic_key("A/1") != genetic_key("A_1")
    assert genetic_key("251-1") != genetic_key("251-80")
    assert genetic_key("é") == genetic_key("e\u0301")


def test_duplicate_input_requires_resolution_without_merging_physical_objects():
    rows = number_rows(GeneticNumbersRequest(examination_year=2026, items=[{"raw_name": "251-1"}, {"raw_name": "251–1"}]))
    assert len(rows) == 2
    assert all(row["error"] for row in rows)


def test_full_sample_key_matching_has_no_fuzzy_examination_match():
    objects = [{"id": 1, "object_no_normalized": "251-1"}, {"id": 2, "object_no_normalized": "251-80"}]
    assert match_genetic_sample("251-80", objects)[:2] == ("matched", 2)
    assert match_genetic_sample("251", objects)[0] == "unmatched"
    aliases = [{"id": 1, "object_no_normalized": "a/1"}, {"id": 2, "object_no_normalized": "a_1"}]
    assert match_genetic_sample("a/1", aliases)[:2] == ("matched", 1)
    assert match_genetic_sample("a/1", [aliases[0], {"id": 3, "object_no_normalized": "a/1"}])[0] == "ambiguous"


def test_rt_genetic_policy_keeps_arbitrary_names_and_registry_policy_is_unchanged():
    rows = [{"sample_name_raw": "АБ-125/1", "normalized_sample_name": "АБ-125/1", "target_key": target, "quantity_ng_ul": quantity} for target, quantity in (("long", 0.8123), ("small", 0.6382), ("y", 0.0142))]
    assert aggregate_rt_rows(rows) == []
    aggregate = aggregate_rt_rows(rows, policy="genetic")
    assert len(aggregate) == 1
    assert [aggregate[0][f"{target}_quantity"] for target in ("long", "small", "y")] == [0.8123, 0.6382, 0.0142]


@pytest.mark.parametrize("policy", ["genetic", "registry"])
def test_both_sources_round_all_rt_targets_to_four_places(policy):
    rows = [{"sample_name_raw": "251-1", "normalized_sample_name": "251-1", "target_key": target, "quantity_ng_ul": quantity} for target, quantity in (("long", 0.81234567), ("small", 0.63825689), ("y", 0.01429445))]
    result = aggregate_rt_rows(rows, policy=policy)[0]
    assert [result[f"{target}_quantity"] for target in ("long", "small", "y")] == [0.8123, 0.6383, 0.0143]
    assert result["targets"]["small"]["quantity"] == 0.6383
    assert rows[1]["quantity_ng_ul"] == 0.63825689


@pytest.mark.parametrize("value", [None, float("nan"), float("inf"), "invalid"])
def test_rt_rounding_keeps_invalid_values_unavailable(value):
    assert round_rt_quantity(value) is None


@pytest.mark.parametrize("value", [0, -1, float("nan"), float("inf"), None])
def test_invalid_concentration_never_creates_dilution(value):
    result = calculate_dilution(value)
    assert result["available"] is False
    assert result["steps"] == []


def test_typed_target_override_changes_only_selected_object_and_keeps_rt():
    objects = [{"id": index, "source_kind": "genetic_examination", "display_name": f"251-{index}", "rt": {"concentration": 0.7, "small_quantity": 0.7, "long_quantity": 0.5, "y_quantity": 0.02}} for index in (1, 2)]
    original = copy.deepcopy(objects)
    plates = build_protocol_layouts(objects)["pcr"]["plates"]
    kwargs = dict(target_concentration=0.1, source_dna_volume=3, dilution_one_volume=10, threshold=100)
    small = build_protocol_dilutions(objects, plates, **kwargs)
    long = build_protocol_dilutions(objects, plates, target_overrides={"genetic_examination:1": "long"}, **kwargs)
    y = build_protocol_dilutions(objects, plates, target_overrides={"genetic_examination:1": "y"}, **kwargs)
    assert [row["source_concentration"] for row in small] == [0.7, 0.7]
    assert [row["source_concentration"] for row in long] == [0.5, 0.7]
    assert y[0]["source_concentration"] == 0.02
    assert y[0]["steps"] == []
    assert long[1] == small[1]
    assert objects == original


def test_missing_selected_target_does_not_fall_back_to_small():
    objects = [{"id": 1, "rcsme_reg_no": "251-1", "rt": {"small_quantity": 0.7, "long_quantity": None}}]
    rows = build_protocol_dilutions(objects, build_protocol_layouts(objects)["pcr"]["plates"], target_concentration=0.1, source_dna_volume=3, dilution_one_volume=10, threshold=100, target_overrides={"chs:1": "long"})
    assert rows[0]["available"] is False
    assert rows[0]["source_concentration"] is None


def test_plate_adapter_uses_genetic_names_and_typed_identity():
    objects = [{"id": 1, "display_name": "251-80", "source_kind": "genetic_examination"}, {"id": 2, "display_name": "251-2", "source_kind": "genetic_examination"}]
    layout = build_protocol_layouts(objects)
    samples = [well for well in layout["pcr"]["plates"][0]["wells"] if well["kind"] == "sample"]
    assert [well["display_name"] for well in samples] == ["251-2", "251-80"]
    assert samples[0]["source_kind"] == "genetic_examination"
    assert protocol_object_count([{"id": 1, "source_kind": "chs"}, {"id": 1, "source_kind": "genetic_examination"}]) == 2


@pytest.mark.parametrize("count", [64, 88, 89, 95, 96, 120])
def test_multiple_genetic_plates_use_the_shared_pcr_overflow_rules(count):
    objects = [{"id": index, "display_name": f"251-{index}", "source_kind": "genetic_examination"} for index in range(1, count + 1)]
    layouts = build_protocol_layouts(objects)
    plates = layouts["pcr"]["plates"]
    assert sum(plate["sample_count"] for plate in plates) == count
    assert all(plate["sample_count"] <= 88 for plate in plates)
    assert len(plates) == (count + 87) // 88
    for plate in plates:
        assert sum(well["kind"] == "ladder" for well in plate["wells"]) == 6
        assert sum(well["kind"] in {"pc", "nc"} for well in plate["wells"]) == 2


def test_imported_layout_is_preserved_and_collisions_are_rejected():
    objects = [{"id": 1, "display_name": "251-1"}, {"id": 2, "display_name": "251-2"}]
    cells = [ProtocolImportedWell(layout_key="pcr", plate_index=1, well=well, kind="sample", source_id=obj_id) for well, obj_id in (("H12", 1), ("A2", 2))]
    result = apply_imported_layouts(build_protocol_layouts(objects), cells, objects)
    assert [(well["well"], well["object_id"]) for well in result["pcr"]["plates"][0]["wells"] if well["kind"] == "sample"] == [("A2", 2), ("H12", 1)]
    with pytest.raises(HTTPException):
        apply_imported_layouts(result, [*cells, cells[0]], objects)
    with pytest.raises(HTTPException):
        apply_imported_layouts(result, cells[:1], objects)


def test_genetic_request_requires_year_and_forbids_parties():
    with pytest.raises(ValueError):
        ProtocolPreviewRequest(kind="genetic_examination", protocol_date="2026-10-14", protocol_no=1, name="test", object_ids=[1])
    with pytest.raises(ValueError):
        ProtocolPreviewRequest(kind="genetic_examination", examination_year=2026, protocol_date="2026-10-14", protocol_no=1, name="test", object_ids=[1], selection={"party_ids": [1]})


def test_excel_policy_keeps_arbitrary_names_controls_and_separate_pcr_grid(tmp_path):
    from openpyxl import Workbook
    workbook = Workbook()
    source = workbook.active
    source.title = "Единая_плашка_A4"
    source["A1"], source["B2"] = "Единый протокол", "14.10.2026"
    source["A10"], source["B10"], source["C10"] = "A", "251-80", "АБ-125/1"
    source["A11"], source["B11"] = "B", "NC"
    pcr = workbook.create_sheet("PCR")
    pcr["A18"], pcr["B18"] = "A", "L1"
    pcr["A19"], pcr["B19"] = "B", "АБ-125/1"
    pcr["A25"], pcr["M25"] = "H", "251-80"
    path = tmp_path / "protocol.xlsx"
    workbook.save(path)
    parsed = parse_work_protocol_preview(path, policy="genetic")
    assert {obj["sample_name_raw"] for obj in parsed.objects} == {"251-80", "АБ-125/1"}
    assert next(cell for cell in parsed.layouts["pcr"] if cell["sample_name_raw"] == "251-80")["well"] == "H12"
    assert next(cell for cell in parsed.layouts["pcr"] if cell["sample_name_raw"] == "L1")["kind"] == "ladder"
    assert parsed.protocol_date == "2026-10-14"


def test_duplicate_excel_labels_keep_physical_identity_and_allow_manual_correction():
    cell = lambda well: {"plate_index": 1, "well": well, "kind": "sample", "sample_name_raw": "SAME"}
    data = {"layouts": {"source": [cell("A1"), cell("B1")], "pcr": [cell("H12"), cell("G12")]}}
    items = excel_source_items(data)
    assert len(items) == 2
    payload = GeneticExcelApply(examination_year=2026, upload_id="00000000000000000000000000000000.xlsx", items=[{"source_key": item["source_key"], "raw_name": f"EX-{index}", "examination_no": "EX"} for index, item in enumerate(items, 1)])
    resolved = resolve_excel_items(payload, data)
    with pytest.raises(HTTPException):
        excel_cells(payload, data, resolved, [11, 12])
    payload.well_mappings = {"pcr:1:H12": "source:1:B1", "pcr:1:G12": "source:1:A1"}
    wells = excel_cells(payload, data, resolved, [11, 12])
    assert [(well["well"], well["source_id"]) for well in wells] == [("A1", 11), ("B1", 12), ("H12", 12), ("G12", 11)]
    payload.well_mappings["pcr:2:A1"] = "source:1:A1"
    with pytest.raises(HTTPException):
        excel_cells(payload, data, resolved, [11, 12])


def test_excel_selected_stage_fields_never_clear_other_settings():
    stage = {"stage_type": "pcr", "settings": {"import": "test"}, "work_date": "2026-10-14", "kit_name": "Kit", "profile_id": None, "performers": [], "performer_ids": [], "comment": None}
    patch = selected_stage_patch(stage, ["kit", "performers", "comment"])
    assert patch == {"stage_type": "pcr", "settings": {"import": "test"}, "kit_name": "Kit", "profile_id": None}


def test_input_sheet_duplicate_names_are_not_silently_merged(tmp_path):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ввод_объектов"
    sheet.append(["№", "Название объекта", "Примечание"])
    sheet.append([1, "251-1", ""])
    sheet.append([2, "251-1", ""])
    path = tmp_path / "input.xlsx"
    workbook.save(path)
    parsed = parse_work_protocol_preview(path, policy="genetic")
    assert len(excel_source_items(parsed.__dict__)) == 2


def test_sparse_physical_rows_are_bounded_without_affecting_registry_reader(tmp_path):
    from openpyxl import Workbook
    from app.parsers.excel_reader import read_workbook
    workbook = Workbook()
    workbook.active["A1"] = "Start"
    workbook.active["A1000000"] = "Unexpected data"
    path = tmp_path / "sparse.xlsx"
    workbook.save(path)
    assert len(next(iter(read_workbook(path).values()))) == 2
    with pytest.raises(ValueError):
        read_workbook(path, preserve_rows=True)
